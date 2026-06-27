"""
W1.3-v4 generator: pilot RUN+EVAL ALLOWLIST, v4 (supersedes v3).

v4 is an EVAL-TIME ALLOWLIST of doc_ids (consumed at W1.5/W3); it does NOT drive
chunking (W1.4 chunks the full corpus). Carries over v3: 95 docs, <100 cap,
balanced 19/theme, deterministic Article-Code-order fill, Biography-only-in-C.

Audit fixes applied (see selection_notes_v4.txt):
  * BALANCED label (not "representative"); corpus is ~55% theme A.
  * Forced lists RE-DERIVED from the corpus Keyword/s field (casefolded tally):
      - orphaned mislabels corrected: CA060 (ULAS/indigents, NOT indigenous
        peoples), SD035 (FLP Scholars/COVID, NOT climate), CA330 (extrajudicial
        shortcuts/due-process, NOT "killings").
      - AI-governance expanded 8 -> 17 (keyword-supported AI/tech subject docs);
        the full keyword-derived TIER-1 set (31) is reported in the notes, with
        the theme-A overflow flagged for a PM call.
  * Theme = coarse curated-theme (ADR-0011 ID letter), to be re-tagged at W1.7.
  * Biography note reworded (GC### = ID-convention artifact).

Reads the four data/csv/*_curated_normalized.xlsx (utf-8-sig content via openpyxl).
Outputs (utf-8-sig) into reports/pilot-eval subset/. Reproducible / deterministic.
v1/v2/v3 files are NOT touched.
"""
from __future__ import annotations
import openpyxl, os, re, csv, ast, collections

OUT_DIR = "reports/pilot-eval subset"
FROZEN_DATE = "2026-06-27"
T = 19
FILES = {"Column": "columns", "Book": "books", "Speech": "speeches", "Biography": "biography"}
FMT_LETTER = {"C": "Column", "B": "Book", "S": "Speech", "G": "Biography"}
THEMES = list("ABCDE")
FORMATS = ["Column", "Biography", "Book", "Speech"]
THEME_LABEL = {
    "A": "Liberty & Rule of Law", "B": "Prosperity & Economic Philosophy",
    "C": "Biographical & Personal", "D": "FLP Mission / Foundation / Judicial Reform",
    "E": "Signature Current-Events Commentary",
}
# corpus theme distribution (read; stated in notes — NOT representative target)
PADDED = re.compile(r"^[CGBS][A-E]\d{3}$")

# ----- forced sets (re-derived) -----
AI_INCLUDED = ["CA034", "CA242", "CB001", "CD022", "CE011", "SA049", "SE012", "SE013",
               "CA031", "CA095", "CA127", "CD011", "CE018", "CE019", "SE036", "CE078", "CE102"]
ORPH = ["CA242", "CA060", "CA377", "CA330", "CC069", "CE072", "BA027", "BA009", "SA049", "SD035"]
# corrected orphaned topic labels + the rare keyword used as evidence
ORPH_LABEL = {
    "CA242": ("right to be forgotten / data privacy", "right to be forgotten"),
    "CA060": ("free legal aid for indigents (ULAS) [CORRECTED from 'indigenous peoples']",
              "Unified Legal Aid Service Ulas"),
    "CA377": ("same-sex marriage", "same-sex marriage"),
    "CA330": ("extrajudicial shortcuts / post-martial-law due process [CORRECTED from 'extrajudicial killings']",
              "extrajudicial shortcuts critique"),
    "CC069": ("surrogacy / assisted reproduction", "surrogacy"),
    "CE072": ("agrarian reform (CASER expropriation)", "agrarian reform expropriation"),
    "BA027": ("reproductive health / RH law", "Reproductive Health"),
    "BA009": ("plagiarism (Del Castillo / Vinuya dispute)", "Del Castillo / Vinuya plagiarism dispute"),
    "SA049": ("intellectual property / jurisprudential disruption", "Ferdinand Negre Kaufman intellectual property"),
    "SD035": ("FLP Scholars Society founding (COVID DiGi Age) [CORRECTED from 'climate']",
              "FLP Scholars Society founding"),
}
# specific AI/tech-governance terms (word-boundary) used for evidence + reporting
TIER1 = [r"\bartificial intelligence\b", r"\bai\b", r"\bmachine learning\b", r"\balgorithm",
         r"\bdeepfake", r"\bblockchain", r"\bgenetic engineering\b", r"\bnanotech",
         r"\bdata privacy\b", r"\bcybercrime\b", r"\bcyberlibel\b",
         r"\bright to be forgotten\b", r"\be-?governance\b"]


def kws(s):
    s = str(s or "").strip()
    if not s:
        return []
    try:
        v = ast.literal_eval(s)
        if isinstance(v, list):
            return [str(x).strip() for x in v]
    except Exception:
        pass
    return [p.strip() for p in s.split(";") if p.strip()]


# ----- load corpus -----
rec = {}
for fmt, stem in FILES.items():
    wb = openpyxl.load_workbook(f"data/csv/cjp_{stem}_curated_normalized.xlsx", read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    h = [str(x).strip() if x is not None else "" for x in next(it)]
    ci, ki, ti = h.index("Article Code"), h.index("Keyword/s"), h.index("Title")
    for row in it:
        if ci >= len(row) or not row[ci]:
            continue
        m = re.match(r"^([CGBS])([A-E])([0-9O]+)$", str(row[ci]).strip().upper())
        if not m:
            continue
        code = f"{m.group(1)}{m.group(2)}{int(m.group(3).replace('O','0')):03d}"
        rec[code] = {"fmt": fmt, "theme": code[1],
                     "title": (str(row[ti]).strip() if row[ti] else ""),
                     "kw": (str(row[ki]) if row[ki] else "")}
    wb.close()
corpus_total = len(rec)

# corpus-wide casefolded keyword tally (for orphaned evidence counts)
tally = collections.Counter()
for r in rec.values():
    for k in set(x.casefold() for x in kws(r["kw"])):
        tally[k] += 1

corpus_cnt = collections.defaultdict(int)
for r in rec.values():
    corpus_cnt[(r["fmt"], r["theme"])] += 1
corpus_theme = collections.Counter(r["theme"] for r in rec.values())


def ai_terms(code):
    kc = " | ".join(kws(rec[code]["kw"])).casefold()
    out = []
    for p in TIER1:
        if re.search(p, kc):
            out.append(p.replace(r"\b", "").replace("-?", "-"))
    return out


# ----- quotas (largest remainder, floor 1) + forced-cell raises -----
def quotas_for_theme(theme):
    present = [f for f in FORMATS if corpus_cnt[(f, theme)] > 0]
    total = sum(corpus_cnt[(f, theme)] for f in present)
    ideal = {f: T * corpus_cnt[(f, theme)] / total for f in present}
    q = {f: int(ideal[f]) for f in present}
    for f in sorted(present, key=lambda f: (-(ideal[f] - int(ideal[f])), FORMATS.index(f)))[:T - sum(q.values())]:
        q[f] += 1
    for f in present:
        if q[f] == 0:
            q[f] = 1
            donor = max((g for g in present if q[g] > 1), key=lambda g: (q[g], -FORMATS.index(g)))
            q[donor] -= 1
    return present, q


forced = sorted(set(AI_INCLUDED) | set(ORPH))
forced_cell = collections.defaultdict(list)
for c in forced:
    forced_cell[(rec[c]["theme"], rec[c]["fmt"])].append(c)

quota, raise_log = {}, []
for theme in THEMES:
    present, q = quotas_for_theme(theme)
    for f in present:
        fc = len(forced_cell.get((theme, f), []))
        if fc > q[f]:
            raise_log.append(f"Theme {theme}: forced {f}={fc} > quota {q[f]} -> raise to {fc}")
            deficit = fc - q[f]; q[f] = fc
            for _ in range(deficit):
                donor = max((g for g in present if g != f and q[g] > 1), key=lambda g: (q[g], -FORMATS.index(g)))
                q[donor] -= 1
                raise_log.append(f"           -> reduce {donor} to {q[donor]}")
    assert sum(q.values()) == T, (theme, q)
    for f in present:
        quota[(theme, f)] = q[f]

# ----- select: forced first, then ascending-code fill -----
selected = {}
def add(code, rationale, cat=""):
    selected[code] = {"theme": rec[code]["theme"], "format": rec[code]["fmt"],
                      "rationale": rationale, "cat": cat}

for c in forced:
    cats = []
    if c in AI_INCLUDED:
        cats.append("AI-governance")
    if c in ORPH:
        cats.append("orphaned-topic")
    desc = []
    if "AI-governance" in cats:
        terms = ai_terms(c)
        desc.append("AI/tech kw=" + (",".join(terms) if terms else "title-only (CD022)"))
    if "orphaned-topic" in cats:
        lbl, ev = ORPH_LABEL[c]
        desc.append(f"orphaned: {lbl} (kw '{ev}' in {tally.get(ev.casefold(),0)} doc)")
    add(c, "Forced [" + " + ".join(cats) + "]: " + " ; ".join(desc), "+".join(cats))

for theme in THEMES:
    for f in FORMATS:
        if (theme, f) not in quota:
            continue
        pool = sorted(c for c, r in rec.items() if r["theme"] == theme and r["fmt"] == f)
        have = sum(1 for c in selected if selected[c]["theme"] == theme and selected[c]["format"] == f)
        for c in pool:
            if have >= quota[(theme, f)]:
                break
            if c in selected:
                continue
            add(c, f"Representative {THEME_LABEL[theme]} / {f} (ascending-code quota fill)")
            have += 1

# ----- emit -----
os.makedirs(OUT_DIR, exist_ok=True)
sel_rows = sorted(selected.items(), key=lambda kv: (kv[1]["theme"], FORMATS.index(kv[1]["format"]), kv[0]))
with open(os.path.join(OUT_DIR, "pilot_subset_frozen_v4.csv"), "w", encoding="utf-8-sig", newline="") as fh:
    w = csv.writer(fh)
    fh.write(f"# FROZEN pilot subset v4 - frozen as of {FROZEN_DATE}. IMMUTABLE - later changes require v5.\n")
    fh.write("# Supersedes v3 (95) / v2 (125) / v1 (25); v1-v3 retained unchanged.\n")
    fh.write("# 95 docs balanced 19/theme, <100 cap. BALANCED allocation for per-theme eval power\n")
    fh.write("# (NOT corpus-representative; corpus is ~55% theme A).\n")
    fh.write("# EVAL-TIME ALLOWLIST consumed at W1.5/W3 - NOT the chunking input (W1.4 chunks the full corpus).\n")
    w.writerow(["doc_id", "theme", "format", "one_line_rationale"])
    for code, d in sel_rows:
        w.writerow([code, d["theme"], d["format"], d["rationale"]])

sel_cnt = collections.defaultdict(int)
for d in selected.values():
    sel_cnt[(d["format"], d["theme"])] += 1
with open(os.path.join(OUT_DIR, "coverage_table_v4.csv"), "w", encoding="utf-8-sig", newline="") as fh:
    w = csv.writer(fh)
    fh.write("# Coverage table v4: theme (A-E) x format. Cell = # selected docs. 19/theme; 95 total (<100).\n")
    w.writerow(["format"] + THEMES + ["ROW_TOTAL"])
    for f in FORMATS:
        cells = [sel_cnt[(f, t)] for t in THEMES]
        w.writerow([f] + cells + [sum(cells)])
    w.writerow(["COL_TOTAL"] + [sum(sel_cnt[(f, t)] for f in FORMATS) for t in THEMES] + [len(selected)])
    w.writerow([])
    w.writerow(["GAP ANALYSIS (empty subset cell that HAS corpus docs = real gap)"])
    w.writerow(["format", "theme", "corpus_docs", "subset_docs", "status"])
    for f in FORMATS:
        for t in THEMES:
            cd, sd = corpus_cnt[(f, t)], sel_cnt[(f, t)]
            status = "expected-empty (none in corpus)" if cd == 0 else ("GAP" if sd == 0 else "covered")
            w.writerow([f, t, cd, sd, status])

with open(os.path.join(OUT_DIR, "forced_inclusions_v4.csv"), "w", encoding="utf-8-sig", newline="") as fh:
    w = csv.writer(fh)
    w.writerow(["category", "doc_id", "format", "theme", "title", "keyword_evidence", "note"])
    for c in AI_INCLUDED:
        terms = ai_terms(c)
        ev = ",".join(terms) if terms else "(no TIER-1 keyword - title-supported only)"
        w.writerow(["AI-governance", c, rec[c]["fmt"], rec[c]["theme"], rec[c]["title"], ev,
                    "expanded re-derivation" if c not in
                    ["CA034","CA242","CB001","CD022","CE011","SA049","SE012","SE013"] else "carried from v2"])
    for c in ORPH:
        lbl, evk = ORPH_LABEL[c]
        w.writerow(["orphaned-topic", c, rec[c]["fmt"], rec[c]["theme"], rec[c]["title"],
                    f"'{evk}' in {tally.get(evk.casefold(),0)} doc(s)", lbl])

# ----- console: report + acceptance -----
print(f"CORPUS TOTAL READ: {corpus_total}")
print("corpus theme dist:", {t: f"{corpus_theme[t]} ({100*corpus_theme[t]/corpus_total:.0f}%)" for t in THEMES})
print("\nFINAL QUOTAS:")
for t in THEMES:
    print(f"  {t}: " + " ".join(f"{f[:3]}={quota.get((t,f),0)}" for f in FORMATS) + f"  total={sum(quota.get((t,f),0) for f in FORMATS)}")
print("RAISE LOG:", raise_log or "(none beyond v3)")
print("\nSELECTED matrix:")
for f in FORMATS:
    print(f"  {f:10}", [sel_cnt[(f, t)] for t in THEMES], "=", sum(sel_cnt[(f, t)] for t in THEMES))
print("  COL_TOTAL ", [sum(sel_cnt[(f, t)] for f in FORMATS) for t in THEMES], "=", len(selected))

ids = list(selected)
checks = [
    ("total < 100", len(ids) < 100, len(ids)),
    ("19 per theme", all(sum(sel_cnt[(f, t)] for f in FORMATS) == 19 for t in THEMES), None),
    ("all forced present", set(forced) <= set(ids), f"{len(set(forced)&set(ids))}/{len(forced)}"),
    ("5 themes + 4 formats", len({d['theme'] for d in selected.values()}) == 5 and len({d['format'] for d in selected.values()}) == 4, None),
    ("ids unique+well-formed+resolve", len(ids) == len(set(ids)) and all(PADDED.match(c) and c in rec for c in ids), None),
    ("no GAP cells", not [(f, t) for f in FORMATS for t in THEMES if corpus_cnt[(f, t)] > 0 and sel_cnt[(f, t)] == 0], None),
]
print("\nACCEPTANCE:")
ok = True
for n, p, v in checks:
    ok &= p
    print(f"  [{'x' if p else ' '}] {n}" + (f": {v}" if v is not None else ""))
print("ALL PASS" if ok else "*** FAIL ***", "| forced count:", len(forced))
