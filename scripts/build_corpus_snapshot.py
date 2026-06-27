"""
W1.4 Step 2 — pin the corpus so "frozen" is verifiable.

Writes corpus_snapshot.json at repo root, containing:
  header     : {producer, date, source, doc_id_count}
  source_files[]: {filename, sha256, bytes, sheet, row_count}
  docs{}     : doc_id -> row_sha256
               (sha256 over the canonical serialization of the doc's curated
                fields: {column_name: cell_string}, keys sorted, UTF-8)

The per-file hash catches any byte change to the xlsx; the per-doc row hash
catches an edit that keeps the file size/structure but changes a row's content.
scripts/verify_pin.py recomputes both and exits nonzero on any drift.

The xlsx themselves are tracked in git (they are < 4 MB total), so the pinned
corpus travels with the repo.
"""
from __future__ import annotations

import hashlib
import json
import re
import sys
from pathlib import Path

import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))
import config

DATA_CSV = PROJECT_ROOT / "data" / "csv"
SNAPSHOT_PATH = PROJECT_ROOT / "corpus_snapshot.json"

SOURCE_FILES = [
    "cjp_columns_curated_normalized.xlsx",
    "cjp_books_curated_normalized.xlsx",
    "cjp_speeches_curated_normalized.xlsx",
    "cjp_biography_curated_normalized.xlsx",
]
# The 15-field curated schema (Link absent in the biography xlsx — handled).
CURATED_COLUMNS = [
    "Date", "Title", "Article Code", "Link", "Keyword/s", "primary_topics",
    "sub_topics", "signature_phrases", "entities", "stances",
    "notable_anecdotes", "target_audience", "register_markers",
    "decision_framework_signals", "one_paragraph_summary",
]
PADDED_RE = re.compile(config.DOC_ID_REGEX_PADDED)
HEADER = {
    "producer": "Sheena / Cowork+Excel W1.2 normalisation",
    "date": "2026-06-26",
    "source": "data/csv/*_curated_normalized.xlsx",
}


def canon_code(raw) -> str | None:
    if raw is None:
        return None
    code = str(raw).strip().upper().replace(" ", "")
    m = re.match(r"^([CGBS])([A-E])([0-9O]+)$", code)
    if not m:
        return None
    return f"{m.group(1)}{m.group(2)}{int(m.group(3).replace('O', '0')):03d}"


def file_sha256(path: Path) -> tuple[str, int]:
    h = hashlib.sha256()
    data = path.read_bytes()
    h.update(data)
    return h.hexdigest(), len(data)


def row_hash(row_map: dict) -> str:
    """Canonical serialization: keys sorted, UTF-8, then sha256."""
    payload = json.dumps(row_map, sort_keys=True, ensure_ascii=False)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def compute_snapshot() -> dict:
    source_files = []
    docs: dict[str, str] = {}
    dupes: list[str] = []
    for fname in SOURCE_FILES:
        path = DATA_CSV / fname
        sha, nbytes = file_sha256(path)
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(it)]
        idx = {name: i for i, name in enumerate(header)}
        ac_i = idx.get("Article Code")
        rows = 0
        for row in it:
            if ac_i is None or ac_i >= len(row) or row[ac_i] in (None, ""):
                continue
            code = canon_code(row[ac_i])
            if code is None or not PADDED_RE.match(code):
                continue
            rows += 1
            row_map = {
                col: ("" if (col not in idx or idx[col] >= len(row)
                             or row[idx[col]] is None) else str(row[idx[col]]))
                for col in CURATED_COLUMNS
            }
            if code in docs:
                dupes.append(code)
            docs[code] = row_hash(row_map)
        wb.close()
        source_files.append({
            "filename": fname, "sha256": sha, "bytes": nbytes,
            "sheet": ws.title, "row_count": rows,
        })
    header = dict(HEADER)
    header["doc_id_count"] = len(docs)
    if dupes:
        header["duplicate_ids"] = sorted(set(dupes))
    return {"header": header, "source_files": source_files, "docs": docs}


def main() -> int:
    snap = compute_snapshot()
    SNAPSHOT_PATH.write_text(
        json.dumps(snap, ensure_ascii=config.JSON_ENSURE_ASCII, indent=2) + "\n",
        encoding=config.OUTPUT_ENCODING)
    print(f"[snapshot] {SNAPSHOT_PATH.name}: {snap['header']['doc_id_count']} docs, "
          f"{len(snap['source_files'])} source files")
    for sf in snap["source_files"]:
        print(f"   {sf['filename']:42} rows={sf['row_count']:4} "
              f"bytes={sf['bytes']:>8} sha={sf['sha256'][:12]}…")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
