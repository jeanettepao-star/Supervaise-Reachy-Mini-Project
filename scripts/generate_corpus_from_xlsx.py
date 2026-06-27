"""
W1.4 — Regenerate the FULL markdown corpus from the canonical W1.2 normalised
xlsx (data/csv/*_curated_normalized.xlsx), NOT the stale 80-doc tracked CSVs.

Fork of scripts/generate_corpus_files.py: same output layout
(corpus/{type}/{theme_folder}/{id}.md + .json) and conventions, but the
metadata source is the four xlsx and the corpus spans all four formats
(columns, books, speeches, biography ~ 1,089 docs).

For each xlsx row:
  - doc_id = "Article Code", canonicalised + padded to ^[CGBS][A-E]\\d{3}$.
  - .json sidecar: the full curated record. JSON-ish cells are parsed;
    `entities` is preserved AS A JSON OBJECT (people/institutions/cases/...),
    never flattened.
  - .md: YAML frontmatter (id/format/type/theme/theme_label/number/title/date/
    keywords/target_audience/word_count/has_body/source_xlsx) + `# Title` +
    the article body (merged from data/text/<id>.md when present, else the
    one_paragraph_summary as fallback) + a `## Summary` section + a
    `## Notable Anecdotes` section (each anecdote under its own `###` heading so
    the W1.4 chunker can keep anecdotes whole).

Knobs (encoding) come from config.py. Idempotent. Old script + old corpus are
left intact; this writes a fresh tree (the caller move-asides the legacy tree).

Usage:
    python scripts/generate_corpus_from_xlsx.py            # write corpus/
    python scripts/generate_corpus_from_xlsx.py --dry-run  # counts only
"""
from __future__ import annotations

import argparse
import ast
import json
import re
import sys
from pathlib import Path

import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))
import config

CORPUS_ROOT = PROJECT_ROOT / "corpus"
DATA_CSV = PROJECT_ROOT / "data" / "csv"
DATA_TEXT = PROJECT_ROOT / "data" / "text"
REPORTS_DIR = PROJECT_ROOT / "reports"

# format letter -> (type label, type folder, xlsx stem)
FORMATS = {
    "C": ("column", "columns", "cjp_columns_curated_normalized.xlsx"),
    "B": ("book", "books", "cjp_books_curated_normalized.xlsx"),
    "S": ("speech", "speeches", "cjp_speeches_curated_normalized.xlsx"),
    "G": ("biography", "biography", "cjp_biography_curated_normalized.xlsx"),
}
THEME_LABELS = {
    "A": "Liberty and Rule of Law",
    "B": "Prosperity and Economic Philosophy",
    "C": "Biographical and Personal",
    "D": "FLP Mission and Foundation",
    "E": "Signature Current Events Commentary",
}
THEME_FOLDERS = {
    "A": "A_liberty_rule_of_law",
    "B": "B_prosperity_economic_philosophy",
    "C": "C_biographical_personal",
    "D": "D_flp_mission_foundation",
    "E": "E_current_events_commentary",
}
# 15-field curated schema (Link is optional — biography xlsx omits it).
SCALAR_FIELDS = ["Date", "Title", "Article Code", "Link", "one_paragraph_summary"]
JSON_CELL_COLUMNS = [
    "Keyword/s", "primary_topics", "sub_topics", "signature_phrases",
    "entities", "stances", "notable_anecdotes", "target_audience",
    "register_markers", "decision_framework_signals",
]
PADDED_RE = re.compile(config.DOC_ID_REGEX_PADDED)
_HEADER_KEY = re.compile(
    r"^\s*(date|publisher|source|by|link|author|venue|occasion|delivered|"
    r"date delivered|time)\b.*?:", re.IGNORECASE)


def canon_code(raw) -> str | None:
    """Canonicalise an Article Code to ^[CGBS][A-E]\\d{3}$ (pad to 3 digits,
    fix O->0 in the number, upper-case). Returns None if uncoercible."""
    if raw is None:
        return None
    code = str(raw).strip().upper().replace(" ", "")
    m = re.match(r"^([CGBS])([A-E])([0-9O]+)$", code)
    if not m:
        return None
    num = m.group(3).replace("O", "0")
    return f"{m.group(1)}{m.group(2)}{int(num):03d}"


def parse_cell(val):
    """Parse a JSON-ish xlsx cell into Python (list/dict/str). Robust to JSON,
    Python-literals, and plain/semicolon strings."""
    if val is None:
        return None
    s = str(val).strip()
    if not s:
        return None
    for loader in (json.loads, ast.literal_eval):
        try:
            return loader(s)
        except Exception:
            pass
    # plain string: split on ';' if it looks like a list, else keep as text
    if ";" in s and "\n" not in s:
        return [p.strip() for p in s.split(";") if p.strip()]
    return s


def extract_body(code: str, title: str, summary: str) -> tuple[str, bool]:
    """Return (body, has_real_body). Body comes from data/text/<code>.md with
    the title line + Key:value header block stripped; falls back to the curated
    summary when no source file exists."""
    p = DATA_TEXT / f"{code}.md"
    if not p.exists():
        return (summary or "").strip(), False
    raw = p.read_text(encoding=config.FILE_ENCODING)
    title_norm = re.sub(r"[^a-z0-9]", "", (title or "").lower())
    out: list[str] = []
    started = False
    for ln in raw.splitlines():
        if not started:
            s = ln.strip()
            if not s:
                continue
            s_norm = re.sub(r"[^a-z0-9]", "", s.lower())
            # drop a leading title line (allowing #/** markup)
            if (not out and title_norm
                    and s_norm[:24] == title_norm[:24]
                    and len(s_norm) <= len(title_norm) + 12):
                continue
            if _HEADER_KEY.match(s):
                continue
            started = True
        out.append(ln)
    body = "\n".join(out).strip()
    if not body:
        return (summary or "").strip(), False
    return body, True


def _yaml_scalar(v) -> str:
    """Emit a value as JSON (valid YAML for scalars/flow lists), ascii-safe off."""
    return json.dumps(v, ensure_ascii=config.JSON_ENSURE_ASCII)


def build_md(rec: dict, body: str, has_body: bool) -> str:
    fm = {
        "id": rec["id"],
        "format": rec["format"],
        "type": rec["type"],
        "theme": rec["theme"],
        "theme_label": rec["theme_label"],
        "number": rec["number"],
        "title": rec["title"],
        "date": rec["date"],
        "keywords": rec.get("keywords") or [],
        "target_audience": rec.get("target_audience") or [],
        "word_count": len(body.split()),
        "has_body": has_body,
        "source_xlsx": rec["source_xlsx"],
    }
    lines = ["---"]
    for k, v in fm.items():
        lines.append(f"{k}: {_yaml_scalar(v)}")
    lines.append("---")
    parts = ["\n".join(lines), f"# {rec['title']}", body]
    summ = (rec.get("one_paragraph_summary") or "").strip()
    if summ:
        parts.append(f"## Summary\n\n{summ}")
    anec = rec.get("notable_anecdotes") or []
    if isinstance(anec, list) and anec:
        sec = ["## Notable Anecdotes"]
        for i, a in enumerate(anec, 1):
            if isinstance(a, dict):
                head = a.get("title") or a.get("label") or f"Anecdote {i}"
                text = a.get("text") or a.get("summary") or a.get("anecdote") or json.dumps(a, ensure_ascii=False)
            else:
                head, text = f"Anecdote {i}", str(a)
            sec.append(f"### {head}\n\n{text}")
        parts.append("\n\n".join(sec))
    return "\n\n".join(parts).strip() + "\n"


def process_format(letter: str, dry_run: bool, stats: dict) -> list[str]:
    type_label, type_folder, stem = FORMATS[letter]
    wb = openpyxl.load_workbook(DATA_CSV / stem, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it)]
    idx = {name: i for i, name in enumerate(header)}
    written: list[str] = []
    for row in it:
        ac = row[idx["Article Code"]] if "Article Code" in idx else None
        if ac is None or str(ac).strip() == "":
            continue  # trailing/blank row
        code = canon_code(ac)
        if code is None or not PADDED_RE.match(code):
            stats["bad_code"].append(str(ac))
            continue
        if code[0] != letter:
            stats["format_mismatch"].append(code)
        theme = code[1]
        title = str(row[idx["Title"]]).strip() if row[idx.get("Title", -1)] else ""
        date = row[idx["Date"]] if "Date" in idx else None
        date_s = str(date)[:10] if date else ""
        parsed = {col: parse_cell(row[idx[col]]) for col in JSON_CELL_COLUMNS if col in idx}
        rec = {
            "id": code, "format": letter, "type": type_label,
            "theme": theme, "theme_label": THEME_LABELS[theme],
            "number": int(code[2:]), "title": title, "date": date_s,
            "link": (str(row[idx["Link"]]).strip() if "Link" in idx and row[idx["Link"]] else ""),
            "keywords": parsed.get("Keyword/s"),
            "primary_topics": parsed.get("primary_topics"),
            "sub_topics": parsed.get("sub_topics"),
            "signature_phrases": parsed.get("signature_phrases"),
            "entities": parsed.get("entities"),       # JSON OBJECT, preserved as-is
            "stances": parsed.get("stances"),
            "notable_anecdotes": parsed.get("notable_anecdotes"),
            "target_audience": parsed.get("target_audience"),
            "register_markers": parsed.get("register_markers"),
            "decision_framework_signals": parsed.get("decision_framework_signals"),
            "one_paragraph_summary": parsed.get("one_paragraph_summary")
                or (str(row[idx["one_paragraph_summary"]]).strip()
                    if "one_paragraph_summary" in idx and row[idx["one_paragraph_summary"]] else ""),
            "source_xlsx": stem,
        }
        body, has_body = extract_body(code, title, rec["one_paragraph_summary"])
        if not has_body:
            stats["no_body"].append(code)
        written.append(code)
        if dry_run:
            continue
        out_dir = CORPUS_ROOT / type_folder / THEME_FOLDERS[theme]
        out_dir.mkdir(parents=True, exist_ok=True)
        (out_dir / f"{code}.md").write_text(build_md(rec, body, has_body),
                                            encoding=config.OUTPUT_ENCODING)
        (out_dir / f"{code}.json").write_text(
            json.dumps(rec, ensure_ascii=config.JSON_ENSURE_ASCII, indent=2) + "\n",
            encoding=config.OUTPUT_ENCODING)
    wb.close()
    return written


def main(argv=None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args(argv)
    stats = {"bad_code": [], "format_mismatch": [], "no_body": []}
    all_ids: list[str] = []
    for letter in ("C", "B", "S", "G"):
        ids = process_format(letter, args.dry_run, stats)
        print(f"[{FORMATS[letter][0]:9}] {len(ids)} docs")
        all_ids.extend(ids)
    dupes = sorted({i for i in all_ids if all_ids.count(i) > 1})
    print(f"[total] {len(all_ids)} docs ({len(set(all_ids))} unique)")
    print(f"[body ] from data/text: {len(all_ids)-len(stats['no_body'])} | "
          f"summary-fallback (no source body): {len(stats['no_body'])}")
    if stats["bad_code"]:
        print(f"[warn ] uncoercible codes: {len(stats['bad_code'])} {stats['bad_code'][:10]}")
    if stats["format_mismatch"]:
        print(f"[warn ] format-letter mismatch: {stats['format_mismatch'][:10]}")
    if dupes:
        print(f"[warn ] duplicate ids: {dupes[:10]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
