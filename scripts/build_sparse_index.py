"""
W1.6 — build the sparse arm: atomic-phrase dictionary + BM25 over FULL chunk text.

Steps (all in memory; outputs written only at the very end so a failure leaves
NO partial files — on any error every output path is removed and the script
exits nonzero):
  1. Build the atomic-phrase dictionary from the curated xlsx (Keyword/s + the
     entities object). Keywords stay atomic (multi-word = one phrase). Entities
     keep canonical + trailing-parenthetical-stripped variants; `cases` verbatim.
  2. Tokenise every W1.4 chunk with the SHARED tokenizer (app/sparse.tokenize) so
     dictionary phrases become atomic BM25 terms with their own IDF.
  3. BM25Okapi(k1,b from config). Pickle {bm25, chunk_ids, doc_ids}.
  4. Write the phrase dict (provenance + counts) and a pin meta (n_chunks,
     n_phrases by provenance, k1/b, sha256 of curated xlsx + chunk index).

No checkpoint/resume scaffolding — BM25 is fast.

Usage:  python scripts/build_sparse_index.py
"""
from __future__ import annotations

import ast
import glob
import hashlib
import json
import pickle
import re
import sys
from pathlib import Path

import openpyxl
from rank_bm25 import BM25Okapi

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))
import config
sys.path.insert(0, str(PROJECT_ROOT / "app"))
import sparse  # noqa: E402

CHUNKS_JSONL = PROJECT_ROOT / "corpus" / "index" / "chunks.jsonl"
CHUNK_INDEX = PROJECT_ROOT / "corpus" / "index" / "chunk_index.json"
_TRAIL_PAREN = re.compile(r"\s*\([^)]*\)\s*$")
_PRIO = {"entity": 0, "keyword": 1, "case": 2}   # higher wins on duplicate


def sha256(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _parse_entities(cell):
    if cell is None or str(cell).strip() == "":
        return {}
    for loader in (json.loads, ast.literal_eval):
        try:
            v = loader(str(cell))
            return v if isinstance(v, dict) else {}
        except Exception:
            continue
    return {}


def _parse_keywords(cell):
    """Atomic keywords from the Keyword/s cell. NOTE: the W1.6 brief says 'split
    on ;' but the curated field is actually a JSON array (the form W1.3/W1.4
    parse) — splitting on ';' would fuse a whole doc's keywords into one bogus
    phrase. So we parse the list and FALL BACK to ';' only for non-list cells.
    (Flagged in the report, not silently 'fixed'.)"""
    if cell is None or str(cell).strip() == "":
        return []
    s = str(cell).strip()
    for loader in (json.loads, ast.literal_eval):
        try:
            v = loader(s)
            if isinstance(v, list):
                return [str(x) for x in v if str(x).strip()]
        except Exception:
            continue
    return [p for p in s.split(";") if p.strip()]


def build_phrase_dict():
    phrases: dict[str, str] = {}

    def add(raw, prov):
        key = sparse.phrase_key(raw)
        if not key:
            return
        existing = phrases.get(key)
        if existing is None or _PRIO[prov] > _PRIO[existing]:
            phrases[key] = prov

    kw_total = kw_multi = 0
    for path in sorted(glob.glob(str(PROJECT_ROOT / config.CURATED_XLSX_GLOB))):
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(it)]
        idx = {name: i for i, name in enumerate(header)}
        ki = idx.get("Keyword/s")
        ei = idx.get("entities")
        for row in it:
            if ki is not None and ki < len(row) and row[ki]:
                for kw in _parse_keywords(row[ki]):
                    key = sparse.phrase_key(kw)
                    if key:
                        kw_total += 1
                        if " " in key:
                            kw_multi += 1
                        add(kw, "keyword")
            if ei is not None and ei < len(row):
                ent = _parse_entities(row[ei])
                for ekey, vals in ent.items():
                    if not isinstance(vals, (list, tuple)):
                        vals = [vals]
                    is_case = str(ekey).strip().lower() == "cases"
                    for v in vals:
                        if not isinstance(v, str) or not v.strip():
                            continue
                        if is_case:
                            add(v, "case")            # verbatim — no strip variant
                        else:
                            add(v, "entity")          # canonical surface form
                            stripped = _TRAIL_PAREN.sub("", v).strip()
                            if stripped and stripped != v:
                                add(stripped, "entity")
        wb.close()
    return phrases, kw_total, kw_multi


def main() -> int:
    outputs = [config.SPARSE_DICT_PATH, config.SPARSE_INDEX_PATH, config.SPARSE_META_PATH]
    try:
        config.SPARSE_INDEX_PATH.parent.mkdir(parents=True, exist_ok=True)

        # 1. dictionary (in memory)
        phrases, kw_total, kw_multi = build_phrase_dict()
        by_prov = {"keyword": 0, "entity": 0, "case": 0}
        for prov in phrases.values():
            by_prov[prov] += 1
        print(f"[sparse] phrase dict: {len(phrases)} unique "
              f"(keyword={by_prov['keyword']} entity={by_prov['entity']} case={by_prov['case']}); "
              f"keywords multi-word {kw_multi}/{kw_total} "
              f"({100*kw_multi/max(1,kw_total):.0f}%)")

        # prime the shared tokenizer with the in-memory dict (no file yet)
        sparse.prime_phrases(phrases.keys())

        # 2-3. tokenise chunks + BM25
        chunk_ids, doc_ids, corpus = [], [], []
        for line in CHUNKS_JSONL.read_text(encoding=config.FILE_ENCODING).splitlines():
            c = json.loads(line)
            chunk_ids.append(c["chunk_id"])
            doc_ids.append(c["doc_id"])
            corpus.append(sparse.tokenize(c["text"]))
        bm25 = BM25Okapi(corpus, k1=config.BM25_K1, b=config.BM25_B)
        print(f"[sparse] BM25 over {len(chunk_ids)} chunks (k1={config.BM25_K1} b={config.BM25_B})")

        # 4. write all outputs (only now that everything succeeded)
        config.SPARSE_DICT_PATH.write_text(
            json.dumps({
                "n_phrases": len(phrases),
                "provenance_counts": by_prov,
                "keyword_multiword_frac": round(kw_multi / max(1, kw_total), 4),
                "phrases": phrases,
            }, ensure_ascii=config.JSON_ENSURE_ASCII, indent=2) + "\n",
            encoding=config.OUTPUT_ENCODING)
        with open(config.SPARSE_INDEX_PATH, "wb") as fh:
            pickle.dump({"bm25": bm25, "chunk_ids": chunk_ids, "doc_ids": doc_ids}, fh)
        meta = {
            "n_chunks": len(chunk_ids),
            "n_phrases": len(phrases),
            "n_phrases_by_provenance": by_prov,
            "bm25_k1": config.BM25_K1,
            "bm25_b": config.BM25_B,
            "chunk_index_sha256": sha256(CHUNK_INDEX),
            "curated_xlsx_sha256": {Path(p).name: sha256(p)
                                    for p in sorted(glob.glob(str(PROJECT_ROOT / config.CURATED_XLSX_GLOB)))},
        }
        config.SPARSE_META_PATH.write_text(
            json.dumps(meta, ensure_ascii=config.JSON_ENSURE_ASCII, indent=2) + "\n",
            encoding=config.OUTPUT_ENCODING)

        print(f"[sparse] wrote {config.SPARSE_DICT_PATH.name}, "
              f"{config.SPARSE_INDEX_PATH.name} ({config.SPARSE_INDEX_PATH.stat().st_size/1024:.0f} KB), "
              f"{config.SPARSE_META_PATH.name}")
        return 0
    except Exception as e:
        for p in outputs:
            Path(p).unlink(missing_ok=True)   # no partial leftovers
        print(f"[sparse] FAILED: {type(e).__name__}: {e} — removed any partial outputs",
              file=sys.stderr)
        raise


if __name__ == "__main__":
    raise SystemExit(main())
